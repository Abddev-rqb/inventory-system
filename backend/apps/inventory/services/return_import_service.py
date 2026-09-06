from django.contrib.auth import (
    get_user_model,
)

from openpyxl import (
    load_workbook,
)

from apps.inventory.models import (
    Laptop,
    Return,
)
from apps.inventory.roles import (
    ROLE_TECHNICIAN,
    get_user_role,
)


User = get_user_model()


class ReturnImportError(
    Exception
):
    pass


class ReturnExcelImportPreviewService:
    SHEET_NAME = (
        "Return Import"
    )

    MAX_ROWS = 5000

    COLUMN_MAP = {
        "customername":
            "customer_name",

        "customeraddress":
            "customer_address",

        "company":
            "company",

        "displaytype":
            "display_type",

        "modelnumber":
            "model_number",

        "processor":
            "processor",

        "processorgeneration":
            "processor_generation",

        "ramgb":
            "ram_gb",

        "storagegb":
            "storage_gb",

        "storagetype":
            "storage_type",

        "serialnumber":
            "serial_number",

        "technician":
            "technician",

        "issue":
            "issue",

        "priority":
            "priority",

        "status":
            "status",

        "servicerack":
            "service_rack",
    }

    REQUIRED_FIELDS = {
        "customer_name",
        "customer_address",
        "company",
        "display_type",
        "model_number",
        "processor",
        "ram_gb",
        "storage_gb",
        "storage_type",
        "serial_number",
        "issue",
        "service_rack",
    }

    ALLOWED_IMPORT_STATUSES = {
        Return.Status.RECEIVED,
        Return.Status.IN_SERVICE,
        Return.Status.REPAIR_COMPLETED,
        Return.Status.SWAP_REQUESTED,
        Return.Status.READY_FOR_DISPATCH,
    }

    @classmethod
    def preview(
        cls,
        uploaded_file,
    ):
        try:
            workbook = (
                load_workbook(
                    uploaded_file,
                    read_only=True,
                    data_only=True,
                )
            )
        except Exception as exc:
            raise ReturnImportError(
                (
                    "The uploaded file "
                    "is not a valid Excel "
                    "workbook."
                )
            ) from exc

        if (
            cls.SHEET_NAME
            not in workbook.sheetnames
        ):
            raise ReturnImportError(
                (
                    "The workbook must "
                    "contain a sheet named "
                    f"'{cls.SHEET_NAME}'."
                )
            )

        worksheet = (
            workbook[
                cls.SHEET_NAME
            ]
        )

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            raise ReturnImportError(
                (
                    "The Return Import "
                    "sheet is empty."
                )
            )

        headers = rows[0]

        field_columns = {}
        ignored_columns = []

        for (
            index,
            raw_header,
        ) in enumerate(
            headers
        ):
            normalized = (
                cls._normalize_header(
                    raw_header
                )
            )

            field_name = (
                cls.COLUMN_MAP.get(
                    normalized
                )
            )

            if field_name:
                field_columns[
                    field_name
                ] = index

            elif raw_header:
                ignored_columns.append(
                    str(
                        raw_header
                    )
                )

        missing_fields = (
            cls.REQUIRED_FIELDS
            - set(
                field_columns
            )
        )

        if missing_fields:
            raise ReturnImportError(
                (
                    "Missing required "
                    "columns: "
                    + ", ".join(
                        sorted(
                            missing_fields
                        )
                    )
                )
            )

        data_rows = rows[1:]

        if (
            len(data_rows)
            > cls.MAX_ROWS
        ):
            raise ReturnImportError(
                (
                    "A maximum of "
                    f"{cls.MAX_ROWS} "
                    "rows can be previewed."
                )
            )

        valid_rows = []
        invalid_rows = []

        seen_serials = set()

        for (
            row_index,
            row,
        ) in enumerate(
            data_rows,
            start=2,
        ):
            if cls._is_empty_row(
                row
            ):
                continue

            raw_data = {}

            for (
                field_name,
                column_index,
            ) in (
                field_columns.items()
            ):
                raw_data[
                    field_name
                ] = (
                    row[
                        column_index
                    ]
                    if column_index
                    < len(row)
                    else None
                )

            (
                cleaned_data,
                errors,
            ) = cls.validate_row(
                raw_data,
                seen_serials=(
                    seen_serials
                ),
            )

            result = {
                "row_number":
                    row_index,

                "data":
                    cleaned_data,
            }

            if errors:
                result[
                    "errors"
                ] = errors

                invalid_rows.append(
                    result
                )

            else:
                seen_serials.add(
                    cleaned_data[
                        "serial_number"
                    ]
                    .lower()
                )

                valid_rows.append(
                    result
                )

        return {
            "total_rows": (
                len(
                    valid_rows
                )
                + len(
                    invalid_rows
                )
            ),

            "valid_rows":
                len(
                    valid_rows
                ),

            "invalid_rows":
                len(
                    invalid_rows
                ),

            "ignored_columns":
                ignored_columns,

            "valid_data":
                valid_rows,

            "invalid_data":
                invalid_rows,
        }

    @classmethod
    def validate_row(
        cls,
        raw_data,
        *,
        seen_serials=None,
    ):
        seen_serials = (
            seen_serials
            if seen_serials
            is not None
            else set()
        )

        errors = []

        customer_name = (
            cls._text(
                raw_data.get(
                    "customer_name"
                )
            )
        )

        customer_address = (
            cls._text(
                raw_data.get(
                    "customer_address"
                )
            )
        )

        company = cls._text(
            raw_data.get(
                "company"
            )
        )

        display_type = (
            cls._normalize_choice(
                raw_data.get(
                    "display_type"
                )
            )
        )

        model_number = (
            cls._text(
                raw_data.get(
                    "model_number"
                )
            )
        )

        processor = cls._text(
            raw_data.get(
                "processor"
            )
        )

        processor_generation = (
            cls._text(
                raw_data.get(
                    "processor_generation"
                )
            )
        )

        ram_gb = (
            cls._positive_integer(
                raw_data.get(
                    "ram_gb"
                )
            )
        )

        storage_gb = (
            cls._positive_integer(
                raw_data.get(
                    "storage_gb"
                )
            )
        )

        storage_type = (
            cls._normalize_choice(
                raw_data.get(
                    "storage_type"
                )
            )
        )

        serial_number = (
            cls._text(
                raw_data.get(
                    "serial_number"
                )
            )
        )

        technician_username = (
            cls._text(
                raw_data.get(
                    "technician"
                )
            )
        )

        issue = cls._text(
            raw_data.get(
                "issue"
            )
        )

        priority = (
            cls._normalize_choice(
                raw_data.get(
                    "priority"
                )
            )
            or Return.Priority.NORMAL
        )

        status = (
            cls._normalize_choice(
                raw_data.get(
                    "status"
                )
            )
            or Return.Status.RECEIVED
        )

        service_rack = cls._text(
            raw_data.get(
                "service_rack"
            )
        )

        required_text = {
            "customer_name":
                customer_name,

            "customer_address":
                customer_address,

            "company":
                company,

            "model_number":
                model_number,

            "processor":
                processor,

            "serial_number":
                serial_number,

            "issue":
                issue,

            "service_rack":
                service_rack,
        }

        for (
            field_name,
            value,
        ) in required_text.items():
            if not value:
                errors.append(
                    {
                        "field":
                            field_name,

                        "message":
                            (
                                "This field "
                                "is required."
                            ),
                    }
                )

        valid_display_types = {
            value
            for value, _
            in Laptop
            .DisplayType
            .choices
        }

        if (
            display_type
            not in valid_display_types
        ):
            errors.append(
                {
                    "field":
                        "display_type",

                    "message":
                        (
                            "Display type must "
                            "be touch or "
                            "non_touch."
                        ),
                }
            )

        valid_storage_types = {
            value
            for value, _
            in Laptop
            .StorageType
            .choices
        }

        if (
            storage_type
            not in valid_storage_types
        ):
            errors.append(
                {
                    "field":
                        "storage_type",

                    "message":
                        (
                            "Storage type must "
                            "be ssd or hdd."
                        ),
                }
            )

        if ram_gb is None:
            errors.append(
                {
                    "field":
                        "ram_gb",

                    "message":
                        (
                            "RAM GB must be "
                            "a positive integer."
                        ),
                }
            )

        if storage_gb is None:
            errors.append(
                {
                    "field":
                        "storage_gb",

                    "message":
                        (
                            "Storage GB must "
                            "be a positive integer."
                        ),
                }
            )

        valid_priorities = {
            value
            for value, _
            in Return
            .Priority
            .choices
        }

        if (
            priority
            not in valid_priorities
        ):
            errors.append(
                {
                    "field":
                        "priority",

                    "message":
                        (
                            "Priority must be "
                            "low, normal, high "
                            "or urgent."
                        ),
                }
            )

        if (
            status
            not in cls
            .ALLOWED_IMPORT_STATUSES
        ):
            errors.append(
                {
                    "field":
                        "status",

                    "message":
                        (
                            "Unsupported return "
                            "status for import."
                        ),
                }
            )

        technician_id = None
        technician_name = ""

        if technician_username:
            try:
                technician = (
                    User.objects.get(
                        username__iexact=(
                            technician_username
                        )
                    )
                )

                if (
                    get_user_role(
                        technician
                    )
                    != ROLE_TECHNICIAN
                ):
                    errors.append(
                        {
                            "field":
                                "technician",

                            "message":
                                (
                                    "Selected user "
                                    "is not a "
                                    "Technician."
                                ),
                        }
                    )

                elif (
                    not technician
                    .is_active
                ):
                    errors.append(
                        {
                            "field":
                                "technician",

                            "message":
                                (
                                    "Technician "
                                    "is inactive."
                                ),
                        }
                    )

                else:
                    technician_id = (
                        technician.id
                    )

                    technician_name = (
                        technician.username
                    )

            except User.DoesNotExist:
                errors.append(
                    {
                        "field":
                            "technician",

                        "message":
                            (
                                "Technician "
                                "username does "
                                "not exist."
                            ),
                    }
                )

        if serial_number:
            serial_key = (
                serial_number.lower()
            )

            if (
                serial_key
                in seen_serials
            ):
                errors.append(
                    {
                        "field":
                            "serial_number",

                        "message":
                            (
                                "Duplicate serial "
                                "number in the "
                                "spreadsheet."
                            ),
                    }
                )

        cleaned_data = {
            "customer_name":
                customer_name,

            "customer_address":
                customer_address,

            "company":
                company,

            "display_type":
                display_type,

            "model_number":
                model_number,

            "processor":
                processor,

            "processor_generation":
                processor_generation,

            "ram_gb":
                ram_gb,

            "storage_gb":
                storage_gb,

            "storage_type":
                storage_type,

            "serial_number":
                serial_number,

            "technician":
                technician_name,

            "technician_username":
                technician_name,

            "issue":
                issue,

            "priority":
                priority,

            "status":
                status,

            "service_rack":
                service_rack,
        }

        return (
            cleaned_data,
            errors,
        )

    @staticmethod
    def _normalize_header(
        value,
    ):
        if value is None:
            return ""

        return (
            str(value)
            .strip()
            .lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )

    @staticmethod
    def _normalize_choice(
        value,
    ):
        if value is None:
            return ""

        return (
            str(value)
            .strip()
            .lower()
            .replace("-", "_")
            .replace(" ", "_")
        )

    @staticmethod
    def _text(
        value,
    ):
        if value is None:
            return ""

        return (
            str(value)
            .strip()
        )

    @staticmethod
    def _positive_integer(
        value,
    ):
        try:
            number = int(
                value
            )
        except (
            TypeError,
            ValueError,
        ):
            return None

        if number < 1:
            return None

        return number

    @staticmethod
    def _is_empty_row(
        row,
    ):
        return not any(
            value not in (
                None,
                "",
            )
            for value in row
        )
