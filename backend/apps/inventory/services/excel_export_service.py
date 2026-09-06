from io import BytesIO

import pandas as pd
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.inventory.models import Laptop


class LaptopExcelExportService:
    SHEET_NAME = "Laptops"

    EXPORT_HEADERS = (
        "ID",
        "Company",
        "Display Type",
        "Model Number",
        "Processor",
        "Processor Generation",
        "RAM GB",
        "Storage GB",
        "Storage Type",
        "Serial Number",
        "Wholesale Price",
        "Retail Price",
        "QC Status",
        "Inventory Status",
        "Comments",
        "Quantity",
        "Warranty Days",
        "Area",
        "Created At",
        "Updated At",
    )

    PRICE_HEADERS = (
        "Wholesale Price",
        "Retail Price",
    )

    DATE_HEADERS = (
        "Created At",
        "Updated At",
    )

    FORMULA_PREFIXES = (
        "=",
        "+",
        "-",
        "@",
    )

    def export(self, queryset=None):
        if queryset is None:
            queryset = Laptop.objects.all()

        export_rows = self._build_export_rows(
            queryset=queryset,
        )

        dataframe = pd.DataFrame(
            export_rows,
            columns=self.EXPORT_HEADERS,
        )

        output = BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl",
        ) as writer:
            dataframe.to_excel(
                writer,
                sheet_name=self.SHEET_NAME,
                index=False,
            )

            worksheet = writer.sheets[
                self.SHEET_NAME
            ]

            self._format_worksheet(
                worksheet=worksheet,
            )

        output.seek(0)

        return {
            "file": output,
            "file_name": self._build_file_name(),
            "exported_rows": len(export_rows),
        }

    def _build_export_rows(self, queryset):
        export_rows = []

        queryset = queryset.iterator(
            chunk_size=1000,
        )

        for laptop in queryset:
            export_rows.append(
                {
                    "ID": laptop.id,
                    "Company": self._safe_text(
                        laptop.company
                    ),
                    "Display Type": (
                        laptop.get_display_type_display()
                    ),
                    "Model Number": self._safe_text(
                        laptop.model_number
                    ),
                    "Processor": self._safe_text(
                        laptop.processor
                    ),
                    "Processor Generation": (
                        self._safe_text(
                            laptop.processor_generation
                        )
                    ),
                    "RAM GB": laptop.ram_gb,
                    "Storage GB": laptop.storage_gb,
                    "Storage Type": (
                        laptop.get_storage_type_display()
                    ),
                    "Serial Number": self._safe_text(
                        laptop.serial_number
                    ),
                    "Wholesale Price": float(
                        laptop.wholesale_price
                    ),
                    "Retail Price": float(
                        laptop.retail_price
                    ),
                    "QC Status": (
                        laptop.get_qc_status_display()
                    ),
                    "Inventory Status": (
                        laptop.get_inventory_status_display()
                    ),
                    "Comments": self._safe_text(
                        laptop.comments
                    ),
                    "Quantity": laptop.quantity,
                    "Warranty Days": (
                        laptop.get_warranty_days_display()
                    ),
                    "Area": self._safe_text(
                        laptop.area
                    ),
                    "Created At": self._excel_datetime(
                        laptop.created_at
                    ),
                    "Updated At": self._excel_datetime(
                        laptop.updated_at
                    ),
                }
            )

        return export_rows

    def _format_worksheet(self, worksheet):
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.row_dimensions[1].height = 24

        self._apply_column_widths(
            worksheet=worksheet,
        )

        self._apply_price_formatting(
            worksheet=worksheet,
        )

        self._apply_date_formatting(
            worksheet=worksheet,
        )

        self._apply_cell_alignment(
            worksheet=worksheet,
        )

    def _apply_column_widths(self, worksheet):
        for column_cells in worksheet.iter_cols():
            column_number = column_cells[0].column
            column_letter = get_column_letter(
                column_number
            )

            maximum_length = 0

            for cell in column_cells:
                if cell.value is None:
                    continue

                cell_length = len(
                    str(cell.value)
                )

                maximum_length = max(
                    maximum_length,
                    cell_length,
                )

            calculated_width = maximum_length + 2

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(calculated_width, 10),
                40,
            )

    def _apply_price_formatting(self, worksheet):
        header_positions = self._header_positions(
            worksheet=worksheet,
        )

        for header in self.PRICE_HEADERS:
            column_number = header_positions[header]

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).number_format = "#,##0.00"

    def _apply_date_formatting(self, worksheet):
        header_positions = self._header_positions(
            worksheet=worksheet,
        )

        for header in self.DATE_HEADERS:
            column_number = header_positions[header]

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).number_format = "yyyy-mm-dd hh:mm:ss"

    def _apply_cell_alignment(self, worksheet):
        for row in worksheet.iter_rows(
            min_row=2,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    @staticmethod
    def _header_positions(worksheet):
        return {
            cell.value: cell.column
            for cell in worksheet[1]
        }

    def _safe_text(self, value):
        if value is None:
            return ""

        text_value = str(value)

        possible_formula = text_value.lstrip()

        if possible_formula.startswith(
            self.FORMULA_PREFIXES
        ):
            return "'" + text_value

        return text_value

    @staticmethod
    def _excel_datetime(value):
        if value is None:
            return None

        if timezone.is_aware(value):
            value = timezone.localtime(value)

        return value.replace(
            tzinfo=None,
        )

    @staticmethod
    def _build_file_name():
        export_date = timezone.localdate().isoformat()

        return (
            f"laptop-inventory-{export_date}.xlsx"
        )